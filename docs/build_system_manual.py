from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r"H:\jpi-football-site\docs\JPi_Football_Intelligence_OS_V17_系统使用说明书.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
TEAL = "137C78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667785"
GOLD = "8A6500"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "1A1A1A"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)
sec.different_first_page_header_footer = True

def font(run, size=11, bold=None, color=BLACK, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    return run

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [
    ("Title", 28, NAVY, 0, 8),
    ("Subtitle", 14, DARK_BLUE, 0, 8),
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = name != "Subtitle"
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for style_name in ["List Bullet", "List Number"]:
    st = styles[style_name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(11)
    st.paragraph_format.left_indent = Inches(0.375)
    st.paragraph_format.first_line_indent = Inches(-0.188)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.25

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def fixed_table(table, widths, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW"); tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total)); tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd"); tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120"); tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row_i, row in enumerate(table.rows):
        if row_i == 0 and header:
            tr_pr = row._tr.get_or_add_trPr(); tbl_header = OxmlElement("w:tblHeader"); tbl_header.set(qn("w:val"), "true"); tr_pr.append(tbl_header)
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr(); tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None: tc_w = OxmlElement("w:tcW"); tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i])); tc_w.set(qn("w:type"), "dxa")
            if row_i == 0 and header: set_cell_shading(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                for r in p.runs: font(r, 9.5, bold=(row_i == 0), color=NAVY if row_i == 0 else BLACK)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        p = t.rows[0].cells[i].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(str(h)), 9.5, True, NAVY)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            font(p.add_run(str(value)), 9.5)
    fixed_table(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)

def bottom_border(paragraph, color="D4DCE3", size="6"):
    p_pr = paragraph._p.get_or_add_pPr(); p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None: p_bdr = OxmlElement("w:pBdr"); p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom"); bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), size); bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), color); p_bdr.append(bottom)

header = sec.header
hp = header.paragraphs[0]
font(hp.add_run("JPi FOOTBALL INTELLIGENCE OS V17.0"), 8.5, True, MID_GRAY)
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
bottom_border(hp)
footer = sec.footer
fp = footer.paragraphs[0]
font(fp.add_run("系统使用说明书  |  第 "), 8.5, color=MID_GRAY)
add_page_number(fp)
font(fp.add_run(" 页"), 8.5, color=MID_GRAY)
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def p(text="", bold_prefix=None, italic=False, color=BLACK, align=None, after=6):
    para = doc.add_paragraph()
    if align is not None: para.alignment = align
    para.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        font(para.add_run(bold_prefix), 11, True, color)
        font(para.add_run(text[len(bold_prefix):]), 11, color=color, italic=italic)
    else: font(para.add_run(text), 11, color=color, italic=italic)
    return para

def bullets(items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        font(para.add_run(item), 11)

def steps(items):
    for item in items:
        para = doc.add_paragraph(style="List Number")
        font(para.add_run(item), 11)

def callout(title, text, kind="info"):
    colors = {"info": (LIGHT_BLUE, DARK_BLUE), "warn": ("FFF4D6", GOLD), "risk": ("FCE8E8", RED), "ok": ("E6F4F1", TEAL)}
    fill, ink = colors[kind]
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    cell = t.cell(0, 0); set_cell_shading(cell, fill)
    cell.paragraphs[0].clear()
    r = cell.paragraphs[0].add_run(f"{title}  "); font(r, 10.5, True, ink)
    r = cell.paragraphs[0].add_run(text); font(r, 10.5, color=BLACK)
    fixed_table(t, [9360], header=False)
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def h1(text): doc.add_paragraph(text, style="Heading 1")
def h2(text): doc.add_paragraph(text, style="Heading 2")
def h3(text): doc.add_paragraph(text, style="Heading 3")
def page_break(): doc.add_page_break()

# Cover
for _ in range(5): doc.add_paragraph()
k = doc.add_paragraph(); k.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(k.add_run("SYSTEM MANUAL  |  V17.0"), 10.5, True, TEAL)
title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_before = Pt(18); title.paragraph_format.space_after = Pt(8)
font(title.add_run("JPi Football Intelligence OS"), 28, True, NAVY)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(sub.add_run("专业足球智能分析系统使用说明书"), 16, False, DARK_BLUE)
desc = doc.add_paragraph(); desc.alignment = WD_ALIGN_PARAGRAPH.CENTER; desc.paragraph_format.space_before = Pt(16)
font(desc.add_run("从数据质量、赛前分析、盘口审计到回测复盘的完整操作指南"), 11, False, MID_GRAY)
for _ in range(6): doc.add_paragraph()
meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(meta.add_run("文档版本 1.0  |  系统版本 OS V17.0\n2026-08-22  |  内部使用"), 10, False, MID_GRAY)
page_break()

# Document control & TOC
h1("文档说明")
table(["项目", "内容"], [
    ["适用对象", "系统管理员、数据维护人员、赛事分析人员、回测复盘人员"],
    ["覆盖范围", "本地开发环境、Sites 私有生产站点、D1 数据库、分析与回测流程"],
    ["当前状态", "网站可运行；球员统计已导入本地 D1；回测库已建立；API-Sports 账户当前暂停"],
    ["安全边界", "API Key 仅允许保存于服务端环境变量；不得写入前端、日志、文档或 Git"],
], [2160, 7200])
callout("重要声明", "系统输出用于数据分析和风险识别，不构成投注或财务建议。任何数据不足的结果必须标记为 Pre-Analysis。", "risk")
h2("目录")
toc = [
    "1. 五分钟快速上手", "2. 系统定位与规则完整性", "3. 安装、启动与访问", "4. 主界面与导航",
    "5. 数据源、同步与质量门", "6. 单场比赛完整分析流程", "7. 输出解读与风险控制", "8. 球员统计数据库",
    "9. 回测复盘数据库", "10. API 与系统接口", "11. 日常操作 SOP", "12. 构建、部署与运维",
    "13. 故障排查", "14. 安全、合规与数据治理", "附录 A. 命令速查", "附录 B. 状态与字段定义"
]
for item in toc:
    para = doc.add_paragraph(style="List Bullet"); font(para.add_run(item), 10.5)
page_break()

h1("1. 五分钟快速上手")
callout("最短路径", "启动本地服务 -> 打开市场总览 -> 刷新数据 -> 选择比赛 -> 检查数据质量 -> 阅读证据融合 -> 仅在规则完整时使用正式结论。", "ok")
h2("1.1 本地启动")
steps([
    "安装 Node.js 22.13.0 或更高版本。",
    "在 PowerShell 中进入 H:\\jpi-football-site。",
    "首次运行执行 npm install；依赖较多，Cloudflare/Sites 工具链可能需要数分钟。",
    "设置 WRANGLER_LOG_PATH 后执行 npx vite --host 0.0.0.0 --port 5173。",
    "浏览器打开 http://localhost:5173/，确认状态栏显示系统在线。",
])
h2("1.2 一场比赛的正确阅读顺序")
steps([
    "先看数据质量等级与时效提示，不从赔率或球队名气直接下结论。",
    "查看实力、核心球员、战意和轮换是否同向。",
    "比较理论盘口与实际盘口，再检查 Bet365 欧亚分歧。",
    "用泊松模型校验进球分布，不把比分模型当作主证据。",
    "确认平局或冷门预警、反证据和风险证据。",
    "最后读取方向、风险和投注价值三类结论，三者不可混为一谈。",
])
callout("停止条件", "数据质量为 C、关键阵容未确认、规则节点缺失或结构冲突无法解释时，结果必须降级为 Pre-Analysis 或 D 级回避。", "warn")

h1("2. 系统定位与规则完整性")
p("JPi Football Intelligence OS V17.0 的目标不是给出一个看似确定的比分，而是识别真实风险、盘口偏离、平局猎物、冷门信号和长期可验证的价值方向。")
h2("2.1 三类结论必须分开")
table(["结论类型", "回答的问题", "常见输出"], [
    ["方向判断", "更可能向哪一侧发展？", "主胜倾向、客队不败、小球倾向"],
    ["风险判断", "结论有多可靠、哪里可能失效？", "A/B/C/D、平局预警、阵容时效风险"],
    ["投注价值判断", "当前价格是否值得承担风险？", "去水概率、EV、Edge、WATCH/正式候选"],
], [1800, 3600, 3960])
h2("2.2 Rule Integrity Principle")
p("任何正式分析必须执行全部已验证有效规则。只要有一个必需节点未执行、未输出或证据不足，分析自动降级为“预分析（Pre-Analysis）”，不得标记为正式分析。")
bullets([
    "强队深盘必须有实力、战意、阵容、泊松四重支撑。",
    "平局只能捕捉，不能为了赔率价值强推。",
    "友谊赛不作为核心投注对象。",
    "比分最多给出 2-3 个，不输出 Top 5。",
    "赛后复盘必须反向校验模型，不允许为赛前判断找借口。",
])

h1("3. 安装、启动与访问")
h2("3.1 环境要求")
table(["组件", "要求", "用途"], [
    ["Node.js", ">= 22.13.0", "运行 Vite/vinext 与 Worker"],
    ["npm", "随 Node.js 安装", "安装锁定依赖"],
    ["浏览器", "Chrome/Edge 最新版", "访问本地仪表盘"],
    ["网络", "可访问数据源", "竞彩/API-Football/公共赛程同步"],
], [1700, 2200, 5460])
h2("3.2 在线与本地地址")
table(["环境", "地址", "说明"], [
    ["Sites 私有生产", "https://jpi-football-intelligence-v17.luonan.chatgpt.site", "需要授权账号登录"],
    ["本地开发", "http://localhost:5173/", "必须先启动 Vite"],
    ["局域网", "终端输出的 Network 地址", "同网段设备可访问，需注意防火墙"],
], [1800, 4300, 3260])
h2("3.3 PowerShell 启动命令")
callout("命令", "Set-Location H:\\jpi-football-site\n$env:WRANGLER_LOG_PATH='.wrangler/wrangler.log'\nnpx vite --host 0.0.0.0 --port 5173", "info")
p("如果出现 ERR_CONNECTION_REFUSED，含义是端口没有服务监听，不是网页地址错误。重新启动 Vite 后再刷新浏览器。")

h1("4. 主界面与导航")
table(["导航项", "主要用途", "关键动作"], [
    ["市场总览", "查看今日比赛、同步状态与风险矩阵", "筛选风险、地区、搜索球队"],
    ["比赛分析", "查看选中比赛的完整证据链", "核对实力、阵容、战意、模型"],
    ["盘口雷达", "按 PDI 和盘口偏离排序", "识别理论盘与实际盘分歧"],
    ["平局猎物", "过滤平局相关信号", "只在多条件共振时触发"],
    ["投注价值库", "观察 EV、Edge 与缺失证据", "WATCH 不等于正式推荐"],
    ["联赛数据库", "查看联赛风格与历史统计", "检查样本、赛季阶段和更新时间"],
    ["回测中心", "查看赛前判断与赛果复盘", "比赛结算后生成样本"],
    ["模型设置", "查看生产锁定参数", "避免浏览器端任意改参"],
], [1700, 4500, 3160])
h2("4.1 顶部状态栏")
bullets([
    "DATA QUALITY：当前数据源质量或实时状态。",
    "ACTIVE MATCHES：已采集比赛数量。",
    "HIGH RISK：D 级结构混乱或数据不足的比赛数。",
    "LAST SYNC：最近一次同步时间；竞彩与临场数据可能存在 2-8 小时滞后。",
    "刷新数据：手动触发同步；页面默认每 5 分钟刷新一次。",
])

h1("5. 数据源、同步与质量门")
h2("5.1 数据源优先级")
table(["优先级", "数据源", "角色", "当前状态"], [
    ["100", "中国体育彩票", "竞彩赛程与官方赔率基准", "已接入"],
    ["90", "API-Football", "Bet365 赛前欧赔与亚洲盘口主源", "账户 Suspended，暂不可用"],
    ["70", "SportsGameOdds", "Bet365 赔率备用源", "需单独密钥"],
    ["50", "ESPN", "国际赛程降级源", "公共数据"],
    ["40", "TheSportsDB", "赛程二级降级源", "公共数据"],
], [1000, 1900, 3900, 2560])
callout("密钥规则", "API_FOOTBALL_KEY 只能写入 .env.local 或 Sites 服务端环境变量。当前 API-Sports 控制台显示 Suspended，接口会返回 Missing application key；恢复账户后再验证。", "warn")
h2("5.2 Data Quality Gate")
table(["等级", "典型条件", "允许输出"], [
    ["A", "首发、伤停、战意、赔率与赛程信息完整，来源可信且时效明确", "可进入正式分析"],
    ["B", "关键数据基本齐全，少量非核心字段缺失", "可分析但保留风险提示"],
    ["C", "阵容或盘口信息明显不足、来源时效偏弱", "只能 Pre-Analysis"],
    ["D", "结构冲突、数据源不可用或核心证据缺失", "不建议参与"],
], [900, 5900, 2560])
h2("5.3 同步故障的安全行为")
bullets([
    "数据源失败不得清空或覆盖最后有效快照。",
    "实时源不可用时显示回退数据并明确标注 FALLBACK DATA。",
    "无赔率时不计算正式 EV；无多公司样本时不推断地区资金流。",
    "数据缺失应降低评级，而不是用历史交锋或球队名气补齐结论。",
])

page_break()
h1("6. 单场比赛完整分析流程")
p("正式分析按下列顺序运行。证据融合优先级为：球队实力 > 核心球员 > 战意 > 数据质量 > 理论盘口 > Bet365 分歧 > 亚盘路径 > 泊松模型 > 热度资金 > 历史交锋。")
table(["序号", "引擎", "必须检查", "主要产出"], [
    ["01", "Data Quality Gate", "首发、伤停、战意、欧亚盘、大小球、赛程", "A/B/C"],
    ["02", "Team Strength", "联赛强度、状态、主客场、攻防效率", "实力层结论"],
    ["03", "Player Impact", "核心首发与关键位置缺失", "阵容修正"],
    ["04", "MDI 战意", "争冠、保级、出线、可接受平局", "战意指数"],
    ["05", "SRS 轮换", "密集赛程、连续客场、下一场关键战", "体能/轮换风险"],
    ["06", "联赛球风", "节奏、对抗、逼抢、转换", "相克修正"],
    ["07", "理论盘口", "去水概率、实力差、主场与联赛修正", "理论让步"],
    ["08", "亚盘审计", "理论盘与实际盘深浅", "保守/合理/偏深"],
], [650, 1900, 4470, 2340])
table(["序号", "引擎", "必须检查", "主要产出"], [
    ["09", "Bet365 分歧", "欧亚深浅、降赔升盘、赔付配合", "平局/冷门预警"],
    ["10", "Poisson", "预期进球与比分概率矩阵", "大小球与 Top 2-3 比分"],
    ["11", "Draw Hunting", "实力接近、平赔稳定、低总进球", "平局预警"],
    ["12", "Evidence Fusion", "主证据、反证据、风险证据", "综合判断"],
    ["13", "Risk Control", "冲突、缺失、时效、热度", "A/B/C/D"],
    ["14", "Output", "方向、风险、价值分离", "结构化结论"],
    ["15", "BackTest", "赛前记录、赛果、命中、误差", "复盘样本"],
    ["16", "Learning", "规则误差与参数更新审批", "学习提案"],
], [650, 1900, 4470, 2340])
h2("6.1 理论盘口与实盘关系")
table(["关系", "解释", "风险提示"], [
    ["理论盘 > 实盘", "机构相对保守", "强方可能支撑不足，检查阵容和战意"],
    ["理论盘 = 实盘", "结构基本合理", "继续看路径、赔付与临场信息"],
    ["理论盘 < 实盘", "实际盘口偏深", "警惕诱上或强压，必须找四重支撑"],
], [1800, 4000, 3560])
h2("6.2 平局猎物启动条件")
bullets([
    "双方实力接近，主队没有形成明确层级优势。",
    "盘口结构不坚决，平赔稳定，降赔不升盘。",
    "战意允许平局，客队防守能力较强。",
    "大小球不支持大胜，泊松集中于 1-1、0-0 或 2-2。",
    "信号不足时必须输出“无平局预警”，不得为了赔率硬推。",
])

h1("7. 输出解读与风险控制")
h2("7.1 单场标准输出")
table(["区块", "读取重点"], [
    ["数据质量", "等级、缺失项、数据时间与来源"],
    ["实力层", "主队优势、客队优势或接近"],
    ["战意层", "必须赢、可接受平局、轮换优先级"],
    ["阵容层", "核心缺阵、伤停、首发和关键位置"],
    ["盘口层", "欧赔、亚盘、理论盘与实际盘一致性"],
    ["Bet365 分歧", "是否触发欧亚分歧预警"],
    ["泊松层", "xG、大小球、最多 2-3 个比分"],
    ["平局预警", "无、轻度、中度或强预警"],
    ["风险评级", "A/B/C/D 及停止条件"],
    ["结论", "胜平负、让球、大小球、比分和建议"],
], [2000, 7360])
h2("7.2 风险等级")
table(["评级", "含义", "操作原则"], [
    ["A", "低风险、证据高度同向", "仍需控制仓位与临场复核"],
    ["B", "中风险、主证据较强但存在反证", "轻仓或组合保护"],
    ["C", "高风险、数据或结构存在明显缺陷", "仅观察，不输出强结论"],
    ["D", "结构混乱或证据冲突无法解释", "回避"],
], [900, 5000, 3460])
callout("价值库提示", "WATCH 表示观察对象，不等于正式投注候选。当前缺少“三家公司、500 场校准、30 分钟报价时效、收盘价 CLV”等门槛时，不生成正式 A/B/C 价值评级。", "warn")

h1("8. 球员统计数据库")
h2("8.1 已导入数据")
table(["指标", "结果"], [
    ["原始 CSV 行数", "87"], ["去重后记录", "85"], ["唯一球员", "83"], ["球队", "42"], ["联赛", "8"], ["赛季范围", "2018-2020"], ["字段数", "59 个源字段 + 来源审计字段"],
], [2800, 6560])
p("文件名包含“current seasons”，但实际数据为 2018-2020。系统保留该时效风险，不将其视为当前阵容或伤停证据。")
h2("8.2 表与主键")
bullets([
    "目标表：player_statistics。",
    "唯一键：player_id + team_id + league_id + league_season。",
    "重复策略：完全相同记录去重；同一唯一键但字段冲突时停止导入并报错。",
    "审计字段：source_file、source_row、imported_at。",
    "重复执行使用 upsert，记录数保持 85，迁移具备幂等性。",
])
h2("8.3 查询接口")
table(["参数", "示例", "说明"], [
    ["limit", "?limit=50", "返回 1-200 条"],
    ["playerId", "?playerId=141", "按 API-Football 球员 ID"],
    ["teamId", "?teamId=13", "按球队 ID"],
    ["leagueId", "?leagueId=1", "按联赛 ID"],
    ["season", "?season=2018", "按赛季"],
    ["search", "?search=Senegal", "球员、球队或联赛模糊搜索"],
], [1500, 2600, 5260])
callout("接口", "GET /api/players/statistics?limit=100", "info")

h1("9. 回测复盘数据库")
h2("9.1 数据结构")
table(["表", "用途", "禁止事项"], [
    ["backtest_runs", "记录一次回测的范围、模型版本、状态和样本量", "不得混合不同模型版本"],
    ["backtest_cases", "逐场保存赛前预测、实际赛果、命中和误差", "不得赛后重写赛前证据"],
    ["backtest_metrics", "按联赛、风险、数据质量等维度聚合指标", "不得隐藏低表现分组"],
    ["learning_updates", "保存规则/参数更新提案及审批状态", "不得自动应用未经审批的改参"],
], [2100, 4250, 3010])
h2("9.2 每场复盘必须记录")
bullets([
    "模型版本、分析状态、数据质量和风险等级。",
    "胜平负、让球、大小球和 2-3 个赛前比分。",
    "主证据、反证据、风险证据及其时间戳。",
    "实际比分、实际赛果和结算来源。",
    "胜平负/让球/大小球是否命中，比分距离、Brier Score、Log Loss 与盈亏单位。",
    "误差类别和可复核的错误详情，不允许只写“运气差”。",
])
h2("9.3 误差分类建议")
table(["类别", "典型问题", "后续动作"], [
    ["data_quality", "首发/伤停/赔率时效错误", "降低来源权重，修正质量门"],
    ["player_impact", "核心缺阵影响估计不足", "更新球员影响参数"],
    ["motivation", "战意或赛制理解错误", "更新战意数据库"],
    ["rotation", "密集赛程和轮换漏判", "修正 SRS"],
    ["handicap", "理论盘与实际盘映射偏差", "更新 PDI/联赛修正"],
    ["poisson", "xG 或比分分布失真", "重新校准联赛进球参数"],
    ["randomness", "红牌、点球、极端效率", "记录但避免过拟合"],
], [1800, 4100, 3460])
h2("9.4 回测概览接口")
callout("接口", "GET /api/backtest/overview。新建数据库初始样本为 0；只有在赛前预测和实际赛果均可追溯时才允许生成回测记录。", "info")

h1("10. API 与系统接口")
table(["方法/路径", "作用", "缓存/注意事项"], [
    ["GET /api/football/live", "竞彩赛程和分析初筛", "优先官方竞彩；失败时降级"],
    ["GET /api/odds/market", "Bet365 欧赔与亚洲盘快照", "API-Football 账户需有效"],
    ["GET /api/leagues", "联赛数据库", "支持 name 筛选"],
    ["GET /api/teams", "球队排名与实力", "无权威数据时返回 seed_fallback"],
    ["GET /api/players/dynamics", "伤停与球员动态", "无数据时不伪造记录"],
    ["GET /api/players/statistics", "导入的历史球员统计", "支持多条件查询"],
    ["GET /api/value-library", "价值库规则与观察状态", "当前 WATCH_ONLY"],
    ["GET /api/backtest/overview", "回测运行、指标与错误分布", "no-store"],
    ["GET /api/data-sources", "数据源注册与健康状态", "密钥仅显示变量名"],
    ["GET /api/system/audit", "规则完整性审计", "不缓存"],
], [2800, 3660, 2900])

h1("11. 日常操作 SOP")
h2("11.1 每日开盘前")
steps([
    "确认网站、D1 和核心数据源在线。",
    "检查 LAST SYNC 与数据源健康状态。",
    "筛除友谊赛、无战意比赛和 D 级比赛。",
    "更新联赛阶段、积分形势、核心伤停和赛程间隔。",
    "运行理论盘口、亚盘审计与欧亚分歧检查。",
])
h2("11.2 临场前")
steps([
    "确认首发和关键球员状态。",
    "重新检查盘口路径、赔付配合和大小球。",
    "核对非临场判断的 2-8 小时滞后风险。",
    "若证据链变化，保留旧版本并生成新分析，不覆盖历史。",
    "只输出 2-3 个比分，并保留反证与停止条件。",
])
h2("11.3 赛后")
steps([
    "从权威来源写入实际比分与赛果。",
    "结算胜平负、让球、大小球和比分距离。",
    "按误差类别复盘，不用赛后结果倒推赛前必然。",
    "更新聚合指标，检查各联赛、风险级别和数据质量分组。",
    "学习更新只生成提案，经审批后应用，并保留可回滚记录。",
])

h1("12. 构建、部署与运维")
h2("12.1 构建")
table(["环境", "命令", "说明"], [
    ["Windows", "npx vinext build", "直接运行五阶段生产构建"],
    ["Bash/Linux", "npm run build", "调用 scripts/build-verified.sh"],
    ["本地开发", "npx vite --host 0.0.0.0 --port 5173", "启动 Worker 与前端"],
], [1600, 4200, 3560])
h2("12.2 Sites 发布顺序")
steps([
    "完成本地构建和接口验证。",
    "提交精确 Git 状态并推送 Sites 源码仓库。",
    "保存 Sites 版本，commit_sha 必须对应推送状态。",
    "使用私有部署发布；若站点共享或公开，必须获得明确批准。",
    "轮询部署状态到 succeeded，并记录生产 URL 与版本号。",
])
callout("生产数据授权", "将本地 CSV 数据推送到 Sites 属于数据外发，必须得到用户明确授权。当前球员统计和回测迁移已在本地完成，但尚未在本说明书生成时自动推送生产。", "risk")
h2("12.3 日志位置")
bullets([
    "本地开发标准输出：.wrangler/local-dev.stdout.log。",
    "本地开发错误输出：.wrangler/local-dev.stderr.log。",
    "Wrangler 状态：.wrangler/state/v3。",
    "本地环境变量：.env.local（Git 忽略）。",
])

h1("13. 故障排查")
table(["症状", "原因", "处理"], [
    ["ERR_CONNECTION_REFUSED", "5173 无监听服务", "启动 Vite；检查 Get-NetTCPConnection -LocalPort 5173"],
    ["vite not recognized", "依赖未完整安装", "npm install 后重新启动"],
    ["API Missing application key", "密钥无效、未激活或账户暂停", "检查 API-Sports 状态与渠道"],
    ["Bet365 bookmaker id unavailable", "目录为空、权限或密钥问题", "先请求 /status，再检查订阅"],
    ["实时源不可用", "网络或上游错误", "保留快照并标记 FALLBACK"],
    ["bash not recognized", "Windows 无 Bash", "改用 npx vinext build"],
    ["D1 binding not found", "D1 由 Vite 动态配置", "本地使用 Miniflare 实际 SQLite 或启动 Worker"],
    ["Sites missing dist/server/index.js", "构建产物不是 vinext/Sites 格式", "使用官方 Sites Vite 适配层"],
], [2400, 3100, 3860])
h2("13.1 快速健康检查")
steps([
    "访问 /api/football/live，确认返回 matches 和 updatedAt。",
    "访问 /api/players/statistics?limit=1，确认 summary.records=85。",
    "访问 /api/backtest/overview，确认接口返回且初始样本为 0。",
    "查看 stderr 日志，不把非阻塞的预打包警告误判为页面失败。",
    "最后请求首页并确认 HTTP 200 与 JPi 标识。",
])

h1("14. 安全、合规与数据治理")
bullets([
    "密钥、令牌、登录信息不得出现在浏览器代码、Git、文档截图或错误日志中。",
    "用户在聊天中公开发送过的密钥建议轮换；Suspended 密钥不应推送生产。",
    "外部数据导入前检查编码、分隔符、字段、空值、重复键和赛季时效。",
    "生产数据外发、公开部署、共享访问必须取得明确授权。",
    "赛前预测必须不可变并带模型版本；赛后复盘使用关联记录，不覆盖原结论。",
    "学习规则更新必须审批、可追溯、可回滚，禁止自动追随单场结果。",
    "用户界面必须保留“非投注建议”和数据时效风险提示。",
])
callout("长期目标", "系统优化目标是跨赛季、跨联赛的稳定校准和风险控制，不是追求单场神预测或绝对命中。", "ok")

page_break()
h1("附录 A. 命令速查")
table(["任务", "PowerShell 命令"], [
    ["进入项目", "Set-Location H:\\jpi-football-site"],
    ["安装依赖", "npm install"],
    ["启动开发", "$env:WRANGLER_LOG_PATH='.wrangler/wrangler.log'; npx vite --host 0.0.0.0 --port 5173"],
    ["生产构建", "npx vinext build"],
    ["检查端口", "Get-NetTCPConnection -LocalPort 5173 -State Listen"],
    ["导入 CSV", "node scripts/import-player-statistics.mjs <CSV路径> drizzle\\0002_player_statistics.sql"],
    ["Git 状态", "git status --short"],
    ["差异检查", "git diff --check"],
], [2100, 7260])
h1("附录 B. 状态与字段定义")
table(["字段/状态", "定义"], [
    ["Pre-Analysis", "规则或数据不完整的预分析，不得标记为正式分析"],
    ["PDI", "理论盘口与实际盘口偏离/结构指标，必须结合证据链解释"],
    ["MDI", "比赛战意指数"],
    ["SRS", "赛程与轮换系统"],
    ["FGP", "谁先进球的结构判断"],
    ["BTP", "双方进球概率判断"],
    ["EV", "模型概率与当前价格结合后的期望价值"],
    ["Edge", "模型概率相对去水市场概率的差值"],
    ["CLV", "投注价格相对收盘价格的价值变化"],
    ["Brier Score", "概率预测误差；越低通常越好"],
    ["Log Loss", "对错误高置信预测惩罚更重的概率损失"],
    ["score_distance", "预测比分与实际比分的距离指标"],
], [2200, 7160])

p("文档结束。系统功能或数据库模式更新后，应同步修订本说明书的版本号、状态说明和接口清单。", italic=True, color=MID_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

# Keep tables and paragraphs stable across renderers.
for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True

doc.core_properties.title = "JPi Football Intelligence OS V17.0 系统使用说明书"
doc.core_properties.subject = "安装、操作、分析、数据、回测与运维指南"
doc.core_properties.author = "JPi Football Intelligence OS"
doc.core_properties.keywords = "JPi, Football Intelligence, 系统说明书, 回测, D1"
doc.save(OUT)
print(OUT)
